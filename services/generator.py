import re
from pathlib import Path

from fpdf import FPDF

from exceptions import DocumentGenerationError
from utils import safe_path, REPORTS_DIR


def _sanitize(value: str) -> str:
    return re.sub(r'[\x00-\x1f\x7f]', '', str(value))


class GeneratorService:

    @staticmethod
    def generate_dynamic_pptx(ai_data: dict, session_id: str) -> Path:
        try:
            from pptx import Presentation
        except ImportError:
            raise DocumentGenerationError("Critical dependency 'python-pptx' is missing.")
        try:
            Path(REPORTS_DIR).mkdir(parents=True, exist_ok=True)
            target = safe_path(REPORTS_DIR, f'Presentation_{session_id}.pptx')

            prs = Presentation()
            prs.slide_width, prs.slide_height = 12192000, 6858000  # 13.333" x 7.5" in EMU

            # Slide 1: Title
            s1 = prs.slides.add_slide(prs.slide_layouts[0])
            s1.shapes.title.text = _sanitize(ai_data.get('report_title', 'EXECUTIVE REPORT')).upper()
            s1.placeholders[1].text = f'InTimeTec Compliance Node  |  Session: {session_id}'

            # Slide 2: Executive Summary
            s2 = prs.slides.add_slide(prs.slide_layouts[1])
            s2.shapes.title.text = 'EXECUTIVE SUMMARY'
            s2.placeholders[1].text = _sanitize(ai_data.get('ai_summary', 'No summary provided.'))

            # Slide 3: Requirements Matrix
            s3 = prs.slides.add_slide(prs.slide_layouts[1])
            s3.shapes.title.text = 'REQUIREMENTS MATRIX'
            tf = s3.placeholders[1].text_frame
            tf.text = 'Extracted Specifications:'
            for row in ai_data.get('structured_data', [])[:5]:
                p = tf.add_paragraph()
                p.text = f"• {_sanitize(row.get('item', 'Parameter'))}: {_sanitize(row.get('value', '—'))}"
                p.level = 1

            prs.save(str(target))
            return target
        except DocumentGenerationError:
            raise
        except Exception as e:
            raise DocumentGenerationError(f'PPTX generation failed: {e}') from e

    @staticmethod
    def generate_dynamic_pdf(ai_data: dict, session_id: str) -> str:
        try:
            Path(REPORTS_DIR).mkdir(parents=True, exist_ok=True)
            target = safe_path(REPORTS_DIR, f'Report_{session_id}.pdf')

            pdf = FPDF()
            pdf.add_page()

            pdf.set_font('Arial', size=18, style='B')
            pdf.set_text_color(44, 62, 80)
            pdf.cell(0, 15, txt=_sanitize(ai_data.get('report_title', 'EXECUTIVE REPORT')).upper(), ln=True, align='C')
            pdf.ln(3)

            pdf.set_font('Arial', size=10, style='I')
            pdf.set_text_color(127, 140, 141)
            pdf.cell(0, 5, txt=f'Report ID: EXP-{session_id.upper()}', ln=True, align='R')
            pdf.ln(8)

            pdf.set_fill_color(52, 73, 94)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font('Arial', size=11, style='B')
            pdf.cell(100, 10, txt='Parameter / Component', border=1, fill=True)
            pdf.cell(90,  10, txt='Value / Calculation',   border=1, fill=True, align='C')
            pdf.ln()

            pdf.set_text_color(44, 62, 80)
            pdf.set_font('Arial', size=11)
            for row in ai_data.get('structured_data', []):
                pdf.cell(100, 10, txt=_sanitize(row.get('item',  'N/A')), border=1)
                pdf.cell(90,  10, txt=_sanitize(row.get('value', 'N/A')), border=1, align='C')
                pdf.ln()

            pdf.ln(8)
            pdf.set_font('Arial', size=12, style='B')
            pdf.set_text_color(44, 62, 80)
            pdf.cell(0, 10, txt='AI Summary & Analytical Insights:', ln=True)

            pdf.set_font('Arial', size=11, style='I')
            pdf.set_text_color(52, 73, 94)
            pdf.multi_cell(0, 7, txt=_sanitize(ai_data.get('ai_summary', 'No details provided.')))

            pdf.output(str(target))
            return str(target)

        except Exception as e:
            raise DocumentGenerationError(f'PDF generation failed: {e}') from e
